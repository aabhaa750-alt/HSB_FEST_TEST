import openpyxl
from pathlib import Path
from backend.database import get_db_connection, init_db

def seed_events():
    """Seeds the events table from data/events_seed.xlsx if the table is currently empty."""
    init_db()

    data_file = Path(__file__).resolve().parent.parent / "data" / "events_seed.xlsx"
    if not data_file.is_file():
        raise FileNotFoundError(
            f"Required seed file not found at {data_file}. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events table already has rows
    cursor.execute("SELECT COUNT(*) AS count FROM events")
    count = cursor.fetchone()["count"]
    if count > 0:
        print(f"Events table already contains {count} rows. Skipping seed.")
        conn.close()
        return

    # Load spreadsheet and read the Events sheet
    workbook = openpyxl.load_workbook(data_file)
    if "Events" not in workbook.sheetnames:
        raise ValueError(f"'Events' sheet not found in {data_file}. Available sheets: {workbook.sheetnames}")

    sheet = workbook["Events"]

    # Header check
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The 'Events' sheet is empty.")

    header = [str(col).strip().lower() if col is not None else "" for col in rows[0]]
    expected_header = ["id", "title", "category", "description", "event_date", "event_time", "venue", "capacity", "fee", "icon"]
    if header != expected_header:
        print(f"Warning: Spreadsheet header {header} differs slightly from expected {expected_header}. Parsing by position.")

    inserted_count = 0
    for row in rows[1:]:
        if not row or all(cell is None for cell in row):
            continue
        
        event_id = row[0]
        title = str(row[1]).strip() if row[1] is not None else ""
        category = str(row[2]).strip() if row[2] is not None else ""
        description = str(row[3]).strip() if row[3] is not None else ""
        event_date = str(row[4]).strip() if row[4] is not None else ""
        event_time = str(row[5]).strip() if row[5] is not None else ""
        venue = str(row[6]).strip() if row[6] is not None else ""
        capacity = int(row[7]) if row[7] is not None else 0
        fee = int(row[8]) if row[8] is not None else 0
        icon = str(row[9]).strip() if row[9] is not None else ""

        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon))
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events into the events table.")

if __name__ == "__main__":
    seed_events()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    events = cursor.fetchall()
    print(f"\n--- Events in Database ({len(events)} total) ---")
    for ev in events:
        print(f"ID {ev['id']:2d} | {ev['icon']} {ev['title']:<24} | {ev['category']:<10} | {ev['event_date']} {ev['event_time']} | {ev['venue']:<20} | Cap: {ev['capacity']:3d} | Fee: ₹{ev['fee']}")
    conn.close()
